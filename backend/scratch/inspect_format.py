import sys
sys.path.insert(0, 'd:/Lexai/phygitron360')
from backend.modules.lexai.api.sop import _prepare_working_docx
from backend.modules.lexai.services import sop_service
from docx import Document

pdf_path = r'C:\Users\Rishit\Downloads\AI_Robotics_Automation_Master_Roadmap_Level_1_to_Final(2).pdf'
with open(pdf_path, 'rb') as f:
    pdf_bytes = f.read()

working_docx = _prepare_working_docx('test.pdf', pdf_bytes)
out_stream, changes, comp = sop_service.apply_sop_formatting(working_docx, {'mode': 'format_only', 'preset': 'ewandz_standard'})

doc = Document(out_stream)
print('Total paragraphs:', len(doc.paragraphs))
print('Total tables:', len(doc.tables))
print('Total sections:', len(doc.sections))

for i, p in enumerate(doc.paragraphs[:20]):
    runs_info = [(r.text[:30], r.font.name, r.font.size.pt if r.font.size else None, r.bold) for r in p.runs]
    print(f'P{i} [style={p.style.name}]: text="{p.text[:60]}"')
    print(f'    runs: {runs_info}')

sec = doc.sections[0]
hdr = sec.header
print('Header paragraphs count:', len(hdr.paragraphs))
for p in hdr.paragraphs:
    print('  Header P text:', repr(p.text))
    # check images in header
    for rel in hdr.part.rels.values():
        if 'image' in rel.reltype:
            print('  Header contains image rel:', rel.target_ref)

ftr = sec.footer
print('Footer paragraphs count:', len(ftr.paragraphs))
for p in ftr.paragraphs:
    print('  Footer P text:', repr(p.text))
