import pytest
from unittest.mock import MagicMock, patch, call
from backend.modules.deploy.repositories.attendance_repo import AttendanceRepository
from backend.core.email_service_extended import send_bimonthly_report_email
from backend.core.scheduler_jobs import run_bimonthly_report

def test_get_tenant_org_admins_finds_admins():
    repo = AttendanceRepository()
    with patch("backend.modules.deploy.repositories.attendance_repo.get_db_connection") as mock_conn_fn:
        mock_conn = MagicMock()
        mock_cur = MagicMock()
        mock_conn_fn.return_value = mock_conn
        mock_conn.cursor.return_value = mock_cur

        # Mock public.tenants query
        mock_cur.fetchone.side_effect = [
            {"admin_email": "orgadmin@company.com", "company_name": "Acme Corp"}, # public.tenants
            {"employee_code": "EMP001", "name": "Alice Admin"}, # employee lookup
        ]

        # Mock tenant users query
        mock_cur.fetchall.return_value = [
            {
                "username": "orgadmin@company.com",
                "employee_code": "EMP001",
                "name": "Alice Admin",
                "email_id": "orgadmin@company.com"
            }
        ]

        admins = repo.get_tenant_org_admins("tenant_acme")
        assert len(admins) == 1
        assert admins[0]["email"] == "orgadmin@company.com"
        assert admins[0]["name"] == "Alice Admin"
        assert admins[0]["employee_code"] == "EMP001"


def test_send_bimonthly_report_email_org_wide():
    sample_report_data = [
        {
            "name": "John Doe",
            "code": "EMP101",
            "stats": {"present": 10, "absent": 0, "half_day": 0, "leave": 0}
        },
        {
            "name": "Jane Smith",
            "code": "EMP102",
            "stats": {"present": 9, "absent": 1, "half_day": 0, "leave": 0}
        }
    ]

    with patch("backend.core.email_service_extended._get_smtp_config", return_value=(None, None, None, None)), \
         patch("backend.core.email_service_extended._mock_log") as mock_log:
        
        success = send_bimonthly_report_email(
            to_email="admin@company.com",
            manager_name="Alice Admin",
            report_data=sample_report_data,
            period_label="16th - 30th September 2026",
            company_name="Acme Corp",
            is_org_wide=True
        )

        assert success is True
        mock_log.assert_called_once()
        args = mock_log.call_args[0]
        assert args[0] == "admin@company.com"
        assert "Organization Attendance Report — 16th - 30th September 2026" in args[1]
        assert "organization" in args[2]


def test_run_bimonthly_report_sends_org_wide_to_admin_and_team_to_manager():
    # Test on the 15th
    with patch("backend.core.scheduler_jobs.datetime") as mock_dt, \
         patch("backend.core.scheduler_jobs.get_db_connection") as mock_db, \
         patch("backend.core.scheduler_jobs.get_active_tenants", return_value=[("tenant_acme", "Acme Corp")]), \
         patch("backend.core.scheduler_jobs.AttendanceService") as mock_service_cls, \
         patch("backend.core.scheduler_jobs.AttendanceRepository") as mock_repo_cls, \
         patch("backend.core.scheduler_jobs.send_bimonthly_report_email") as mock_send_email:

        mock_dt.now.return_value.year = 2026
        mock_dt.now.return_value.month = 9
        mock_dt.now.return_value.day = 15

        mock_conn = MagicMock()
        mock_cur = MagicMock()
        mock_db.return_value = mock_conn
        mock_conn.cursor.return_value = mock_cur
        mock_cur.fetchone.return_value = None  # Idempotency check: not sent yet

        mock_service = MagicMock()
        mock_service_cls.return_value = mock_service

        mock_repo = MagicMock()
        mock_repo_cls.return_value = mock_repo

        # Org admin
        mock_repo.get_tenant_org_admins.return_value = [
            {"email": "admin@company.com", "name": "Alice Admin", "employee_code": "EMP001"}
        ]

        # Managers (one of whom is the org admin, and another is team lead Bob)
        mock_repo.get_all_managers.return_value = ["EMP001", "EMP002"]
        mock_repo.get_employee_email.side_effect = lambda code, tenant: {
            "EMP001": "admin@company.com",
            "EMP002": "bob@company.com"
        }.get(code)
        mock_repo.get_employee_name.side_effect = lambda code, tenant: {
            "EMP001": "Alice Admin",
            "EMP002": "Bob Lead"
        }.get(code)

        # Mock report data
        org_report_data = [{"name": "All Staff", "code": "EMP999", "stats": {}}]
        bob_team_data = [{"name": "Bob Team Member", "code": "EMP888", "stats": {}}]
        
        mock_service.get_bimonthly_report.side_effect = lambda y, m, c, manager_code: (
            org_report_data if manager_code is None else bob_team_data
        )

        run_bimonthly_report()

        # Should send 2 emails:
        # 1. To admin@company.com with is_org_wide=True
        # 2. To bob@company.com with is_org_wide=False
        # EMP001 (admin) should NOT get a duplicate team email!
        assert mock_send_email.call_count == 2
        
        # Verify 1st call is org-wide to admin
        call1 = mock_send_email.call_args_list[0]
        assert call1.kwargs["to_email"] == "admin@company.com"
        assert call1.kwargs["is_org_wide"] is True
        assert call1.kwargs["report_data"] == org_report_data

        # Verify 2nd call is team-specific to Bob
        call2 = mock_send_email.call_args_list[1]
        assert call2.kwargs["to_email"] == "bob@company.com"
        assert call2.kwargs["is_org_wide"] is False
        assert call2.kwargs["report_data"] == bob_team_data


def test_generate_attendance_excel_structure():
    import io
    import openpyxl
    from backend.common.utils.excel_utils import generate_attendance_excel

    sample_report = [
        {
            "name": "Jane Employee",
            "code": "EMP010",
            "team": "Engineering",
            "designation": "Frontend Dev",
            "days": [
                {"day": 1, "date": "2026-09-16", "status": "Present"},
                {"day": 2, "date": "2026-09-17", "status": "Absent"},
                {"day": 3, "date": "2026-09-18", "status": "Half Day"},
                {"day": 4, "date": "2026-09-19", "status": "Leave"},
                {"day": 5, "date": "2026-09-20", "status": "Weekend"},
            ],
            "stats": {"present": 1, "absent": 1, "half_day": 1, "leave": 1, "holiday": 0}
        }
    ]

    excel_bytes = generate_attendance_excel(
        sample_report,
        period_label="16th - 20th September 2026",
        company_name="Test Corp",
        is_org_wide=True
    )

    assert excel_bytes is not None
    assert len(excel_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet_names = wb.sheetnames
    assert "Attendance Summary" in sheet_names
    assert "Daily Matrix" in sheet_names

    ws_summary = wb["Attendance Summary"]
    assert "Test Corp" in ws_summary["A1"].value
    assert "Organization-Wide" in ws_summary["A1"].value
    assert ws_summary["B5"].value == "Jane Employee"
    assert ws_summary["C5"].value == "EMP010"
    assert ws_summary["E5"].value == 1  # present
    assert ws_summary["F5"].value == 1  # absent

    ws_matrix = wb["Daily Matrix"]
    assert ws_matrix["B3"].value == "Jane Employee"


def test_download_attendance_excel_endpoint_org_wide():
    from fastapi.testclient import TestClient
    from backend.main import app
    from backend.core.dependencies import get_current_user
    from backend.modules.deploy.services.attendance_service import AttendanceService
    import io
    import openpyxl

    client = TestClient(app)

    admin_user = {
        "user_id": 1,
        "username": "admin@test.com",
        "role": "org_admin",
        "roles": ["org_admin"],
        "employee_code": "EMP001",
        "tenant_id": "test_tenant",
        "permissions": {
            "module.deploy.access": True,
            "deploy.attendance.view_all": True,
            "deploy.attendance.view_team": True
        }
    }

    mock_report = [
        {
            "name": "Alice Admin",
            "code": "EMP001",
            "team": "Operations",
            "days": [{"day": 1, "date": "2026-09-16", "status": "Present"}],
            "stats": {"present": 1, "absent": 0, "half_day": 0, "leave": 0, "holiday": 0}
        }
    ]

    app.dependency_overrides[get_current_user] = lambda: admin_user

    try:
        with patch.object(AttendanceService, "get_company_name", return_value="Acme Corp"), \
             patch.object(AttendanceService, "get_attendance_report", return_value=mock_report) as mock_get_rep:

            # 1. Test custom date range download
            resp = client.get("/api/attendance/report/download-excel?start_date=2026-09-16&end_date=2026-09-30&scope=org")
            assert resp.status_code == 200
            assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            assert "Attendance_Report_OrgWide_2026-09-16_to_2026-09-30.xlsx" in resp.headers["content-disposition"]
            mock_get_rep.assert_called_with("2026-09-16", "2026-09-30", None)

            # Ensure valid excel content
            wb = openpyxl.load_workbook(io.BytesIO(resp.content))
            assert "Attendance Summary" in wb.sheetnames

            # 2. Test preset cycle download
            resp_cycle = client.get("/api/attendance/report/download-excel?year=2026&month=9&cycle=1")
            assert resp_cycle.status_code == 200
            assert "Attendance_Report_OrgWide_2026_09_1st_to_15th.xlsx" in resp_cycle.headers["content-disposition"]
            mock_get_rep.assert_called_with("2026-09-01", "2026-09-15", None)
    finally:
        app.dependency_overrides.pop(get_current_user, None)


def test_download_attendance_excel_endpoint_manager_forbidden_for_org_wide():
    from fastapi.testclient import TestClient
    from backend.main import app
    from backend.core.dependencies import get_current_user

    client = TestClient(app)

    manager_user = {
        "user_id": 2,
        "username": "manager@test.com",
        "role": "employee",
        "roles": ["employee"],
        "employee_code": "MGR002",
        "tenant_id": "test_tenant",
        "permissions": {
            "module.deploy.access": True,
            "deploy.attendance.view_team": True
        }
    }

    app.dependency_overrides[get_current_user] = lambda: manager_user

    try:
        resp = client.get("/api/attendance/report/download-excel?start_date=2026-09-16&end_date=2026-09-30&scope=org")
        assert resp.status_code == 403
        assert "Organization-wide report requires deploy.attendance.view_all" in resp.json()["detail"]
    finally:
        app.dependency_overrides.pop(get_current_user, None)

